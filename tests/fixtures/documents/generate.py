"""Regenerate synthetic fixtures (python-docx, reportlab, openpyxl, xlwt)."""
from pathlib import Path
import subprocess
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
import xlwt
root = Path(__file__).parent
text = '内容索引测试 秦前坤 document_needle_42'
doc = Document(); doc.add_paragraph(text)
doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格内文字 table_needle_73'
doc.save(root / 'sample.docx')
source = root / 'source.txt'; source.write_text(text + '\n', encoding='utf-8')
subprocess.run(['/usr/bin/textutil', '-convert', 'doc', '-format', 'txt', '-inputencoding', 'UTF-8', '-output', str(root / 'sample.doc'), str(source)], check=True)
source.unlink()
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
pdf = canvas.Canvas(str(root / 'sample.pdf')); pdf.setFont('STSong-Light', 12)
pdf.drawString(50, 750, text); pdf.save()
pdf = canvas.Canvas(str(root / 'image-only.pdf')); pdf.rect(10, 10, 80, 80, fill=1); pdf.save()
book = Workbook(); book.active.append([text, 123.45]); book.create_sheet('第二张表').append(['table_needle_73'])
book.save(root / 'sample.xlsx')
old = xlwt.Workbook(); old.add_sheet('首张表').write(0, 0, text)
old.add_sheet('第二张表').write(0, 0, 'table_needle_73')
old.save(str(root / 'sample.xls'))
